# Generates a workbook with one sheet per bank-export layout, used to check
# that XlsxImport finds each table and its columns without bank-specific code.
# Synthetic data only.
#
# Run from repo root:  python test/fixtures/files/imports/generate_formats.py
import datetime as dt
from openpyxl import Workbook

D = dt.date


def put(ws, top_left_row, top_left_col, rows):
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is not None:
                ws.cell(row=top_left_row + r, column=top_left_col + c, value=value)


wb = Workbook()
wb.remove(wb.active)

# 1. Plain table at A1.
put(wb.create_sheet("a1_simple"), 1, 1, [
    ["Date", "Description", "Amount"],
    [D(2024, 3, 1), "Coffee shop", -3.5],
    [D(2024, 3, 2), "Salary", 2500],
    [D(2024, 3, 3), "Groceries", -54.2],
])

# 2. Title lines and blank rows above the table.
put(wb.create_sheet("titles"), 1, 1, [
    ["Relevé de compte"],
    ["Titulaire : M. Exemple"],
    ["Période du 01/03/2024 au 31/03/2024"],
    [],
    ["Date opération", "Libellé", "Montant"],
    [D(2024, 3, 1), "PRLV EDF", -80.12],
    [D(2024, 3, 5), "VIR SALAIRE", 2100],
])

# 3. Table that doesn't start in column A.
ws = wb.create_sheet("offset")
ws["A1"] = "Export"
put(ws, 6, 3, [
    ["Date", "Payee", "Amount"],
    [D(2024, 4, 1), "Bakery", -4.1],
    [D(2024, 4, 2), "Refund", 12],
])

# 4. Dates and amounts stored as text (French formats).
put(wb.create_sheet("text_values"), 1, 1, [
    ["Mon relevé"],
    [],
    ["Date", "Libellé", "Montant"],
    ["15/03/2024", "CARTE MONOPRIX", "-1 234,56 €"],
    ["16/03/2024", "VIREMENT RECU", "500,00 €"],
])

# 5. UK style: positive "Paid out" / "Paid in" columns and a balance column.
put(wb.create_sheet("paid_in_out"), 1, 1, [
    ["Current account", "12-34-56 12345678"],
    ["Statement date", D(2024, 5, 31)],
    [],
    ["Date", "Type", "Description", "Paid out", "Paid in", "Balance"],
    [D(2024, 5, 1), "DD", "COUNCIL TAX", 120.0, None, 880.0],
    [D(2024, 5, 2), "BGC", "EMPLOYER LTD", None, 1800.0, 2680.0],
])

# 6. Key/value metadata block holding numbers and dates above the table.
put(wb.create_sheet("key_values"), 1, 1, [
    ["Account number", 12345678],
    ["Opening balance", 1500.0],
    ["Export date", D(2024, 6, 30)],
    [],
    ["Booking date", "Text", "Amount", "Currency"],
    [D(2024, 6, 1), "Rent", -900, "EUR"],
    [D(2024, 6, 3), "Gym", -30, "EUR"],
])

# 7. German export: split Soll / Haben, both positive.
put(wb.create_sheet("german"), 1, 1, [
    ["Kontoumsätze"],
    [],
    ["Buchungstag", "Valuta", "Verwendungszweck", "Soll", "Haben", "Währung"],
    [D(2024, 7, 1), D(2024, 7, 1), "Miete Juli", 750.0, None, "EUR"],
    [D(2024, 7, 2), D(2024, 7, 2), "Gehalt", None, 3000.0, "EUR"],
])

# 8. Totals footer below the data.
put(wb.create_sheet("totals_footer"), 1, 1, [
    ["Date", "Label", "Amount"],
    [D(2024, 8, 1), "Shop", -10],
    [D(2024, 8, 2), "Shop", -20],
    [],
    ["Total", None, -30],
])

# 9. Amount column before the date column.
put(wb.create_sheet("amount_first"), 1, 1, [
    ["Amount", "Date", "Payee", "Category"],
    [-15.99, D(2024, 9, 1), "Streaming", "Leisure"],
    [-42.0, D(2024, 9, 2), "Fuel", "Transport"],
])

# 10. Date-time cells.
put(wb.create_sheet("datetime"), 1, 1, [
    ["Transaction date", "Merchant", "Amount"],
    [dt.datetime(2024, 10, 1, 14, 35), "Taxi", -18.4],
    [dt.datetime(2024, 10, 2, 9, 5), "Cashback", 2.5],
])

# 11. Group header row above the real header.
put(wb.create_sheet("two_line_header"), 1, 1, [
    ["Opération", None, "Montants"],
    ["Date", "Libellé", "Débit", "Crédit"],
    [D(2024, 11, 1), "LOYER", -700.0, None],
    [D(2024, 11, 2), "REMBOURSEMENT", None, 45.0],
])

# 12. Summary block (labels + numbers) above the table.
put(wb.create_sheet("summary_block"), 1, 1, [
    ["Synthèse", "Période"],
    ["-500,00", "800,00"],
    [],
    ["Date", "Libellé", "Montant", "Solde"],
    [D(2024, 12, 1), "ACHAT", -20.0, 980.0],
    [D(2024, 12, 2), "DEPOT", 100.0, 1080.0],
])

# 13. Spanish headers with a balance column.
put(wb.create_sheet("spanish"), 1, 1, [
    ["Movimientos de la cuenta"],
    ["Fecha", "Concepto", "Importe", "Saldo"],
    [D(2024, 1, 10), "Supermercado", -35.1, 964.9],
    [D(2024, 1, 11), "Nómina", 1900.0, 2864.9],
])

# 14. Unrecognised amount header: the only non-date numeric column.
put(wb.create_sheet("unknown_amount_header"), 1, 1, [
    ["Jour", "Opération", "Somme"],
    [D(2024, 2, 1), "Boulangerie", -2.3],
    [D(2024, 2, 2), "Remboursement", 15.0],
])

# 15. Merged title/banner cells: the exporter repeats the merged value in
# every covered cell, including a banner between the header and the data.
put(wb.create_sheet("merged_banners"), 1, 1, [
    ["Encours prélevé : Fin août 2026"] * 3,
    ["R.I.B. : 10278 02619 00099999901 - 08/2026"] * 3,
    [],
    ["Liste de vos comptes"] * 3,
    ["Date", "Libellé", "Montant", "Dev"],
    ["Carte Mastercard XXXXXXXXXXXX9999 (EUR)"] * 4,
    [D(2026, 8, 18), "OPENAI CHATGPT", -23.0, "EUR"],
    [D(2026, 8, 17), "ENTERPRISE RENT", 200.0, "EUR"],
])

# 16-19. Headers in other supported languages and scripts.
put(wb.create_sheet("polish"), 1, 1, [
    ["Historia rachunku"],
    [],
    ["Data operacji", "Opis", "Obciążenia", "Uznania", "Saldo"],
    [D(2024, 3, 4), "Biedronka", -45.2, None, 955.0],
    [D(2024, 3, 5), "Wynagrodzenie", None, 5200.0, 6155.0],
])
put(wb.create_sheet("russian"), 1, 1, [
    ["Выписка по счёту"],
    ["Дата", "Описание", "Сумма", "Остаток"],
    [D(2024, 4, 1), "Пятёрочка", -1250.5, 8749.5],
    [D(2024, 4, 2), "Зарплата", 60000.0, 68749.5],
])
put(wb.create_sheet("turkish"), 1, 1, [
    ["Hesap hareketleri"],
    [],
    ["İşlem tarihi", "Açıklama", "Tutar", "Bakiye"],
    [D(2024, 5, 2), "Market alışverişi", -320.75, 4679.25],
    [D(2024, 5, 3), "Maaş", 25000.0, 29679.25],
])
put(wb.create_sheet("chinese"), 1, 1, [
    ["交易明细"],
    ["交易日期", "摘要", "支出", "收入", "余额"],
    [D(2024, 6, 1), "超市购物", 88.5, None, 911.5],
    [D(2024, 6, 2), "工资", None, 8000.0, 8911.5],
])

# 20. Unsigned amounts with a debit/credit indicator column.
put(wb.create_sheet("dc_indicator"), 1, 1, [
    ["Date", "Description", "Amount", "D/C"],
    [D(2024, 7, 5), "Bakery", 4.2, "D"],
    [D(2024, 7, 6), "Salary", 2100.0, "C"],
])

# 21. Dates written as text with month names, in several languages.
put(wb.create_sheet("month_names"), 1, 1, [
    ["Date", "Libellé", "Montant"],
    ["5 janv. 2024", "Boulangerie", -4.2],
    ["Feb 6, 2024", "Salary", 2100.0],
    ["7 de marzo de 2024", "Supermercado", -35.1],
    ["8. April 2024", "Miete", -750.0],
    ["2024年5月9日", "工资", 8000.0],
])

# 22. Two accounts stacked in one sheet, each with its own title and header.
put(wb.create_sheet("stacked_tables"), 1, 1, [
    ["Compte courant"],
    ["Date", "Libellé", "Montant"],
    [D(2024, 1, 5), "Boulangerie", -4.2],
    [D(2024, 1, 6), "Salaire", 2100.0],
    [],
    ["Livret A"],
    ["Date", "Libellé", "Montant"],
    [D(2024, 2, 1), "Intérêts", 12.5],
    [D(2024, 3, 1), "Versement", 100.0],
])

# 23. Notes only, no table.
put(wb.create_sheet("notes_only"), 1, 1, [
    ["Ce fichier a été généré automatiquement."],
    ["Aucune opération sur la période."],
])

out = "test/fixtures/files/imports/bank_formats.xlsx"
wb.save(out)
print("wrote", out)

# Workbook saved with the 1904 date system (old Mac Excel).
from openpyxl.utils.datetime import CALENDAR_MAC_1904

wb = Workbook()
wb.epoch = CALENDAR_MAC_1904
ws = wb.active
ws.title = "Compte"
put(ws, 1, 1, [["Date", "Libellé", "Montant"], [D(2024, 1, 5), "Boulangerie", -4.2]])
out = "test/fixtures/files/imports/date1904.xlsx"
wb.save(out)
print("wrote", out)
